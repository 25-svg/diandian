from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"C:\Users\10230\Documents\Codex\Workspace\01-Projects\Project-003-直播切片分析系统"
    r"\06-交付物\于千惠直播话术母稿分析与新人培训手册_2026-07-24.docx"
)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_BLUE = "EAF3FF"
LIGHT_GRAY = "F5F7FA"
CAUTION = "FFF7E6"
WHITE = "FFFFFF"


def set_font(run, name="Microsoft YaHei", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=130, bottom=90, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    tbl_pr.append(ind)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(int(width.twips)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def no_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_paragraph_border(paragraph, color=BLUE, size="10", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True)
        add_text(p, text[len(bold_prefix):])
    else:
        add_text(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.16
    add_text(p, text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.16
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    style = doc.styles[f"Heading {level}"]
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size={1: 16, 2: 13, 3: 11.5}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [Inches(6.5)])
    cell = table.cell(0, 0)
    shade(cell, fill)
    no_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_text(p, title, color=DARK_BLUE, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, body, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_metadata(doc):
    rows = [
        ("培训对象", "二手影像直播间新主播、场控与复盘人员"),
        ("话术来源", "于千惠样本批次，6场直播整理"),
        ("当前母稿", "MS-BATCH-2 V1.0，已审核发布"),
        ("适用范围", "金典拍拍二手影像直播间的需求接待、讲品、成交与售后说明"),
        ("使用原则", "先确认当场事实，再使用固定结构；不把价格、库存、赠品或型号当作固定承诺"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [Inches(1.35), Inches(5.15)])
    for row, (label, value) in zip(table.rows, rows):
        shade(row.cells[0], LIGHT_GRAY)
        for cell in row.cells:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
        add_text(row.cells[0].paragraphs[0], label, bold=True, color=DARK_BLUE)
        add_text(row.cells[1].paragraphs[0], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_structure_table(doc):
    headers = ["环节", "新人的动作", "可直接使用的表达", "当场必须确认"]
    rows = [
        ("1. 欢迎与领券", "欢迎新进用户，告诉对方可问型号。", "欢迎老师们光临金典拍拍，店铺券先领一下；想看什么型号可以直接打在公屏上。", "店铺券入口、当日活动"),
        ("2. 需求确认", "先问预算、卡口、用途和定/变焦偏好。", "老师您预算大概多少？配什么卡口？平时主要拍什么，我按您的需求给您找。", "预算、卡口、用途、偏好"),
        ("3. 实物讲解", "拿实物，先讲成色，再讲检测和适用性。", "这只就是您看到的实物，我把外观、卡口、镜片/CMOS、光圈和对焦给您看一遍。", "具体型号、成色、检测结果"),
        ("4. 报价与链接", "说满减或粉丝券后到手价，再说链接位置。", "这只在小黄车X号链接，优惠后到手是XX元；您看中这只我给您按实物备注。", "链接号、优惠、到手价"),
        ("5. 下单促成", "明确拍什么规格，说明赠品与备注动作。", "您拍XX新/带包装这个规格，拍完和我说一声，我给您备注好；直播间下单有当场确认的福利。", "规格、赠品、备注方式"),
        ("6. 售后兜底", "把二手顾虑落到服务事实，不空泛保证。", "咱家支持到手试机；售后按当场页面和客服确认的规则执行，您有问题可直接从订单找到我们。", "退换、保修、物流、运费险规则"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [Inches(0.9), Inches(1.3), Inches(2.9), Inches(1.4)])
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, DARK_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, header, size=9.3, color=WHITE, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_text(p, value, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_advantage_table(doc):
    headers = ["可说明的优势", "样本中的事实依据", "新人标准表达"]
    rows = [
        ("实物一致", "直播间展示实物，并强调“所见即所得”。", "老师，直播间给您看的就是按这只实物来安排，外观和关键部位我都给您看清楚。"),
        ("成色透明", "会展示磨损位置、卡口、镜片/CMOS、光圈和对焦。", "二手不怕有使用痕迹，关键是提前说清；这只哪里有痕迹、哪些功能已确认，我都按实物给您讲。"),
        ("售后可追溯", "样本多次出现：到手试机、7天无理由、1年店铺保修、顺丰和运费险。", "购买后的售后规则以当场页面和客服确认为准；您从订单就能找到我们继续处理。"),
        ("直播间服务", "按需求找货、补库存或调货；可配置店铺券、粉丝券、赠品。", "您先把预算和需求告诉我，我按当场库存给您找；能安排的优惠和赠品，我在下单前一次说明白。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [Inches(1.15), Inches(2.45), Inches(2.9)])
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, DARK_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, header, size=9.4, color=WHITE, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            add_text(p, value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "于千惠直播话术新人培训手册 | 内部培训使用 | ", size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for level, size, color in ((1, 16, BLUE), (2, 13, BLUE), (3, 11.5, DARK_BLUE)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
        style.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, "新人培训手册", size=11, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "于千惠直播话术母稿分析", size=24, color=INK, bold=True)
    set_paragraph_border(p)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    add_text(p, "二手影像直播间：从接待、讲品到下单与售后的可复用表达", size=12, color=MUTED)

    add_callout(
        doc,
        "培训结论",
        "于千惠的话术核心不是硬推，而是“先问清需求，再用实物和售后减少顾虑”。新人应学习成交结构和表达顺序，不应照搬历史价格、库存、赠品、型号或链接。",
    )
    add_metadata(doc)

    add_heading(doc, "一、这套话术的核心方法")
    add_body(doc, "于千惠的直播采用需求型导购：互动收需求、确认预算和卡口、现场讲实物、算到手价、指向小黄车，再用试机和售后规则完成风险兜底。")
    for item in [
        "先让用户说需求：不清楚需求时，先问预算、用途、卡口、定焦或变焦偏好。",
        "先讲实物再成交：展示外观、磨损位置和功能检测，不用模糊的“很好”代替事实。",
        "报价必须完整：链接位置、优惠规则、到手价、对应规格应在同一轮说清。",
        "承诺必须有边界：价格、库存、物流、保修和赠品只按当场商品页、客服或实物确认。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "二、新人应掌握的六步成交结构")
    add_structure_table(doc)

    add_heading(doc, "三、可直接用于直播的标准表达")
    scripts = [
        ("需求确认", "老师您预算大概多少？配什么卡口？平时主要拍什么，我按您的需求给您找，不给您随便乱配。"),
        ("实物讲解", "这只就是您现在看到的实物。我先给您看外观和边角，再看卡口、镜片/CMOS、光圈和对焦；哪里有使用痕迹我提前说清楚。"),
        ("报价下单", "这只我给您放到小黄车X号链接。优惠规则按页面来，您拍下前我再和您确认一次对应型号、成色和到手价。"),
        ("库存与发货", "这只目前的库存和发货时间，我按当场仓库信息给您确认；如果暂时不在手边，我先查仓或给您找合适的替代。"),
        ("售后收尾", "二手买的是实物和售后。您到手先试机，售后规则以商品页面和客服确认的信息为准；后续有问题可以从订单直接找到我们。"),
    ]
    for title, script in scripts:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, f"{title}：", color=DARK_BLUE, bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.18
        add_text(p, f"“{script}”", color=INK)

    add_heading(doc, "四、为什么用户可以选择我们")
    add_body(doc, "样本中已经有“选择我们”的事实基础，但分散在讲品、物流、售后和福利环节。新人不能用贬低同行或绝对化承诺成交，应围绕可证明的服务事实说明优势。")
    add_advantage_table(doc)
    add_callout(
        doc,
        "店铺优势统一说明（确认后使用）",
        "老师，二手买的最重要就是成色和售后。咱直播间给您看的就是按这只实物来安排，外观和关键部位我都提前讲清楚。价格、优惠和库存我按当场页面确认；下单后的物流和售后规则，也以商品页面和客服确认的信息为准。您有问题可以从订单直接找到我们继续处理。",
        fill=CAUTION,
    )

    add_heading(doc, "五、必须遵守的表达边界")
    for item in [
        "不能承诺“全网最低”“永远有货”“一定当天发货”“任何情况都可退”，除非当场规则明确支持。",
        "不能把其他主播、历史场次或其他商品的价格、库存、赠品和链接复制到当前商品。",
        "型号、成色和功能信息不确定时，先看实物、商品页或问仓库；不要把转写错误当成事实。",
        "用户问到竞争商家时，回到自身可验证的服务：实物展示、成色透明、售后规则和订单可追溯。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "六、新人开播前检查清单")
    for item in [
        "确认当前商品的型号、成色、价格、优惠、库存、链接号和赠品。",
        "确认物流时效、运费险、退换及店铺保修规则是否适用于当场商品。",
        "准备需求确认问题：预算、卡口、用途、定焦/变焦、是否需要配件。",
        "讲品时按“实物外观 → 关键部位 → 功能状态 → 价格链接 → 售后边界”顺序走。",
        "用户拍下后，重复确认型号、规格、备注和联系方式，不口头替代订单信息。",
    ]:
        add_number(doc, item)

    add_callout(
        doc,
        "资料说明",
        "本手册基于于千惠样本批次已审核发布的 MS-BATCH-2 V1.0 整理。它用于培训成交结构与表达方式，不替代商品页面、客服政策、仓库库存或实际经营数据。",
        fill=LIGHT_GRAY,
    )

    doc.core_properties.title = "于千惠直播话术母稿分析与新人培训手册"
    doc.core_properties.subject = "二手影像直播新人培训"
    doc.core_properties.author = "典典直播切片"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
